#!/usr/bin/env python3
"""Generate matching web data and PDFs from the controlled SOP DOCX sources."""
from __future__ import annotations

import html
import json
import re
from pathlib import Path

from docx import Document
from reportlab import rl_config
from docx.table import Table
from docx.text.paragraph import Paragraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph as PdfParagraph,
    SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
rl_config.invariant = 1
OUTPUT = ROOT / "generated" / "sops"
SOURCES = (
    ("inventory-lot-control", ROOT / "docs/Inventory_and_Lot_Control_SOP_Rev_03.docx", "ORA-APP-SOP-001", "Rev 03"),
    ("order-corrections-cancellations", ROOT / "deliverables/app-sop-templates/06_Order_Corrections_and_Cancellations_SOP.docx", "ORA-APP-SOP-002", "Rev 01"),
    ("daily-fulfillment-pick-list", ROOT / "deliverables/app-sop-templates/07_Daily_Fulfillment_and_Pick_List_SOP.docx", "ORA-APP-SOP-003", "Rev 03"),
    ("period-end-reporting", ROOT / "deliverables/app-sop-templates/08_Period_End_Reporting_SOP.docx", "ORA-APP-SOP-004", "Rev 01"),
)
REQUIRED_HEADINGS = {
    "1. Purpose and scope", "2. Who performs this", "3. Before you begin",
    "4. Procedure", "5. Problems and exceptions", "6. Records and completion",
    "Revision history",
}


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def clean(value):
    return " ".join((value or "").split())


def validate_control_metadata(
    path, control, approval, banner, expected_id, expected_revision,
    require_banner=True,
):
    expected_banner = "CONTROLLED SOP — DRAFT FOR APPROVAL"
    if require_banner and banner != expected_banner:
        raise SystemExit(f"{path}: required draft banner is missing or changed")
    if banner and banner != expected_banner:
        raise SystemExit(f"{path}: expected draft banner {expected_banner!r}, found {banner!r}")
    if control.get("SOP ID") != expected_id or control.get("Revision") != expected_revision:
        raise SystemExit(
            f"{path}: expected {expected_id} {expected_revision}, found "
            f"{control.get('SOP ID')} {control.get('Revision')}"
        )
    if control.get("Effective date") != "Upon approval":
        raise SystemExit(f"{path}: effective date must remain 'Upon approval'")
    if approval.get("Approved by") != "Quality":
        raise SystemExit(f"{path}: expected Quality approver")
    if approval.get("Approval date") != "Pending":
        raise SystemExit(f"{path}: approval date must remain Pending")
    if approval.get("Document status") != "Editable master — approval copy not yet released":
        raise SystemExit(f"{path}: document status no longer identifies an unreleased approval draft")
    return "DRAFT FOR APPROVAL"


def validate_core_metadata(path, core_properties, expected_revision):
    revision_number = int(expected_revision.removeprefix("Rev ").strip())
    if core_properties.revision != revision_number:
        raise SystemExit(
            f"{path}: expected DOCX core revision {revision_number}, "
            f"found {core_properties.revision!r}"
        )
    if revision_number > 1 and expected_revision not in (core_properties.title or ""):
        raise SystemExit(
            f"{path}: DOCX core title must identify {expected_revision}"
        )


def parse_source(slug, path, expected_id, expected_revision):
    if not path.is_file():
        raise SystemExit(f"Missing SOP source: {path.relative_to(ROOT)}")
    document = Document(path)
    validate_core_metadata(path, document.core_properties, expected_revision)
    title = clean(next((p.text for p in document.paragraphs if p.style.name == "Title"), ""))
    if not title:
        title = clean(next((
            p.text for p in document.paragraphs
            if clean(p.text) and "CONTROLLED SOP" not in p.text.upper()
        ), ""))
    if not title:
        raise SystemExit(f"{path}: missing Title paragraph")
    if len(document.tables) < 2:
        raise SystemExit(f"{path}: missing control metadata tables")

    control = {
        clean(document.tables[0].rows[0].cells[i].text): clean(document.tables[0].rows[1].cells[i].text)
        for i in range(len(document.tables[0].columns))
    }
    approval = {
        clean(document.tables[1].rows[0].cells[i].text): clean(document.tables[1].rows[1].cells[i].text)
        for i in range(len(document.tables[1].columns))
    }
    banner = clean(next((
        p.text for p in document.paragraphs
        if clean(p.text).upper().startswith("CONTROLLED SOP —")
    ), ""))
    status = validate_control_metadata(
        path, control, approval, banner, expected_id, expected_revision,
        require_banner=(slug != "period-end-reporting"),
    )

    blocks, headings = [], set()
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = clean(block.text)
            if not text or block.style.name == "Title":
                continue
            if block.style.name.startswith("Heading"):
                level = int(re.search(r"\d+", block.style.name).group())
                headings.add(text)
                blocks.append({"type": "heading", "level": level, "text": text, "anchor": anchor(text)})
            elif block.style.name.startswith("List Number"):
                blocks.append({"type": "ordered", "text": text})
            elif block.style.name.startswith("List Bullet"):
                blocks.append({"type": "bullet", "text": text})
            else:
                kind = "warning" if text.upper().startswith(("STOP:", "WARNING:", "CAUTION:")) else "paragraph"
                blocks.append({"type": kind, "text": text})
        else:
            rows = [[clean(cell.text) for cell in row.cells] for row in block.rows]
            if rows and any(any(cell for cell in row) for row in rows):
                flattened = " ".join(cell for row in rows for cell in row).strip()
                if len(rows) == 1 and len(rows[0]) == 1 and flattened.upper().startswith(
                    ("STOP:", "WARNING:", "CAUTION:")
                ):
                    blocks.append({"type": "warning", "text": flattened})
                else:
                    blocks.append({"type": "table", "rows": rows})

    missing = REQUIRED_HEADINGS - headings
    if missing:
        raise SystemExit(f"{path}: missing required sections: {', '.join(sorted(missing))}")
    if not any(b["type"] == "ordered" for b in blocks):
        raise SystemExit(f"{path}: no numbered procedure content found")
    searchable = " ".join(
        b.get("text", "") if b["type"] != "table"
        else " ".join(cell for row in b["rows"] for cell in row)
        for b in blocks
    ).upper()
    if not any(b["type"] == "warning" for b in blocks):
        raise SystemExit(f"{path}: no STOP/WARNING/CAUTION content found")

    return {
        "slug": slug, "title": title, "sop_id": expected_id, "revision": expected_revision,
        "status": status, "effective_date": control["Effective date"],
        "approval_status": f"{approval['Approved by']} approval pending",
        "document_status": approval["Document status"],
        "source": str(path.relative_to(ROOT)), "blocks": blocks,
    }


def anchor(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def pdf_text(text):
    return html.escape(text).replace("—", "&mdash;").replace("–", "&ndash;")


def render_pdf(sop, target):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SopTitle", parent=styles["Title"], textColor=colors.HexColor("#1B2A4A"), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="Warning", parent=styles["BodyText"], backColor=colors.HexColor("#FFF4DB"), borderColor=colors.HexColor("#D48A00"), borderWidth=1, borderPadding=8, spaceBefore=8, spaceAfter=8))
    doc = SimpleDocTemplate(str(target), pagesize=letter, rightMargin=.65*inch, leftMargin=.65*inch, topMargin=.65*inch, bottomMargin=.65*inch,
                            title=f"{sop['sop_id']} {sop['revision']} {sop['title']}")
    story = [PdfParagraph(pdf_text(sop["title"]), styles["SopTitle"])]
    metadata = [
        ["SOP ID", sop["sop_id"], "Revision", sop["revision"]],
        ["Status", sop["status"], "Effective date", sop["effective_date"]],
        ["Approval", sop["approval_status"], "Document status", sop["document_status"]],
    ]
    table = PdfTable([[PdfParagraph(pdf_text(c), styles["BodyText"]) for c in row] for row in metadata], colWidths=[.9*inch, 2*inch, 1.05*inch, 2.55*inch])
    table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .5, colors.HexColor("#CBD3E0")), ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EEF2F7")), ("BACKGROUND",(2,0),(2,-1),colors.HexColor("#EEF2F7")), ("VALIGN",(0,0),(-1,-1),"TOP"), ("PADDING",(0,0),(-1,-1),5)]))
    story += [table, Spacer(1, 14)]
    pending = []
    list_kind = None
    for block in sop["blocks"]:
        kind = block["type"]
        if kind in {"ordered", "bullet"}:
            if pending and list_kind != kind:
                story.append(ListFlowable([ListItem(PdfParagraph(pdf_text(x), styles["BodyText"])) for x in pending], bulletType="1" if list_kind == "ordered" else "bullet", leftIndent=24))
                pending = []
            list_kind = kind
            pending.append(block["text"])
            continue
        if pending:
            story.append(ListFlowable([ListItem(PdfParagraph(pdf_text(x), styles["BodyText"])) for x in pending], bulletType="1" if list_kind == "ordered" else "bullet", leftIndent=24))
            pending = []
        if kind == "heading":
            story.append(PdfParagraph(pdf_text(block["text"]), styles["Heading1" if block["level"] == 1 else "Heading2"]))
        elif kind == "warning":
            story.append(PdfParagraph(pdf_text(block["text"]), styles["Warning"]))
        elif kind == "paragraph":
            story.append(PdfParagraph(pdf_text(block["text"]), styles["BodyText"]))
            story.append(Spacer(1, 5))
        elif kind == "table":
            rows = [[PdfParagraph(pdf_text(cell), styles["BodyText"]) for cell in row] for row in block["rows"]]
            widths = [(7.2*inch)/max(1, len(rows[0]))] * len(rows[0])
            t = PdfTable(rows, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.HexColor("#CBD3E0")), ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E8EEF6")), ("VALIGN",(0,0),(-1,-1),"TOP"), ("PADDING",(0,0),(-1,-1),4)]))
            story.append(t)
            story.append(Spacer(1, 8))
    if pending:
        story.append(ListFlowable([ListItem(PdfParagraph(pdf_text(x), styles["BodyText"])) for x in pending], bulletType="1" if list_kind == "ordered" else "bullet", leftIndent=24))
    doc.build(story)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    catalog = []
    for args in SOURCES:
        sop = parse_source(*args)
        json_path = OUTPUT / f"{sop['slug']}.json"
        pdf_path = OUTPUT / f"{sop['slug']}.pdf"
        json_path.write_text(json.dumps(sop, indent=2, ensure_ascii=False) + "\n")
        render_pdf(sop, pdf_path)
        catalog.append({key: sop[key] for key in ("slug", "title", "sop_id", "revision", "status", "effective_date", "approval_status", "document_status")})
    (OUTPUT / "catalog.json").write_text(json.dumps(catalog, indent=2, ensure_ascii=False) + "\n")
    print(f"Published {len(catalog)} SOP drafts to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()